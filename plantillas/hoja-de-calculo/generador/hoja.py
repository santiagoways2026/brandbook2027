# Plantilla de hoja de cálculo Santiago Ways (Brandbook 2027)
# Compatible con Excel y con Hojas de cálculo de Google.
# Uso: python3 hoja.py salida.xlsx
import sys, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.drawing.image import Image as XLImage
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.formatting.rule import FormulaRule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.comments import Comment

OUT = sys.argv[1] if len(sys.argv) > 1 else 'Hoja_Santiago_Ways.xlsx'
D = os.path.dirname(os.path.abspath(__file__))

GRAFITO, GRIS, NIEBLA, WAYS, SENDERO, BROTE, BLANCO = '2E2E2D', '6F6F6E', 'E8E8E6', '7AA606', '506718', 'FAFFEE', 'FFFFFF'
FONT = 'Montserrat'

def font(size=10, bold=False, color=GRAFITO):
    return Font(name=FONT, size=size, bold=bold, color=color)

thin = lambda c: Side(style='thin', color=c)
medium = lambda c: Side(style='medium', color=c)
FILL_IN = PatternFill('solid', fgColor=BROTE)        # celdas para rellenar
EUR = '#,##0 "€";-#,##0 "€";"-"'
INT = '#,##0;-#,##0;"-"'
PCT = '0.0%;-0.0%;"-"'

wb = Workbook()
# fuente por defecto del libro: Montserrat
wb._fonts[0] = Font(name=FONT, size=10, color=GRAFITO)
for st in wb._named_styles:
    if st.name == 'Normal':
        st.font = Font(name=FONT, size=10, color=GRAFITO)


def setup(ws, widths, landscape=True):
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 100
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    ws.page_setup.orientation = 'landscape' if landscape else 'portrait'
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins.left = ws.page_margins.right = 0.5
    ws.page_margins.top = ws.page_margins.bottom = 0.6


def header_block(ws, eyebrow, title, desc, last_col):
    """Cabecera común: logo, antetítulo, título y descripción."""
    img = XLImage(os.path.join(D, 'logo.png'))
    img.width, img.height = 170, 27
    ws.add_image(img, 'B2')
    ws.row_dimensions[1].height = 14
    ws.row_dimensions[2].height = 26
    ws.row_dimensions[3].height = 14
    ws['B4'] = eyebrow.upper(); ws['B4'].font = font(8, True, SENDERO)
    ws['B5'] = title; ws['B5'].font = font(18, True)
    ws.row_dimensions[5].height = 30
    ws['B6'] = desc; ws['B6'].font = font(10, False, GRIS)
    ws.row_dimensions[7].height = 10


def table_header(ws, row, cols, labels, aligns):
    for c, t, a in zip(cols, labels, aligns):
        cell = ws[f'{c}{row}']
        cell.value = t.upper()
        cell.font = font(8, True)
        cell.alignment = Alignment(horizontal=a, vertical='bottom', wrap_text=True)
        cell.border = Border(bottom=medium(GRAFITO))
    ws.row_dimensions[row].height = 30


def body_cell(cell, fmt=None, align='left', fill=None, bold=False):
    cell.font = font(10, bold)
    cell.alignment = Alignment(horizontal=align, vertical='center')
    cell.border = Border(bottom=thin(NIEBLA))
    if fmt:
        cell.number_format = fmt
    if fill:
        cell.fill = fill


def total_cell(cell, fmt=None, align='left'):
    cell.font = font(10, True)
    cell.alignment = Alignment(horizontal=align, vertical='center')
    cell.border = Border(top=medium(GRAFITO))
    if fmt:
        cell.number_format = fmt


# ================= 1 · PORTADA =================
ws = wb.active
ws.title = 'Portada'
setup(ws, {'A': 3, 'B': 22, 'C': 22, 'D': 22, 'E': 22, 'F': 3})
img = XLImage(os.path.join(D, 'logo.png'))
img.width, img.height = 230, 37
ws.add_image(img, 'B2')
ws.row_dimensions[2].height = 34
for r in range(3, 8):
    ws.row_dimensions[r].height = 20
ws['B9'] = 'HOJA DE CÁLCULO · 2027'; ws['B9'].font = font(8, True, SENDERO)
ws['B10'] = 'Título del archivo'; ws['B10'].font = font(24, True)
ws.row_dimensions[10].height = 40
ws['B11'] = 'Descripción breve de para qué sirve esta hoja.'; ws['B11'].font = font(11, False, GRIS)
ws.row_dimensions[12].height = 24
for i, (k, v) in enumerate([('Autor', 'Nombre Apellido'), ('Área', 'Marketing'), ('Fecha', 'dd/mm/aaaa'), ('Versión', '1.0')]):
    c = 'BCDE'[i]
    ws[f'{c}13'] = k.upper(); ws[f'{c}13'].font = font(7, True, GRIS); ws[f'{c}13'].border = Border(top=medium(GRAFITO))
    ws[f'{c}14'] = v; ws[f'{c}14'].font = font(10); ws[f'{c}14'].fill = FILL_IN
ws.row_dimensions[15].height = 30
ws['B16'] = 'CÓMO USAR ESTA PLANTILLA'; ws['B16'].font = font(8, True, SENDERO)
guide = [
    ('Celdas en verde claro', 'Son las que se rellenan o se cambian. El resto se calcula solo.'),
    ('Pestañas de ejemplo', 'Datos, Presupuesto, Plan y Resumen. Duplica la que necesites y borra las demás.'),
    ('Cifras', 'Los datos de ejemplo son inventados: sustitúyelos por los reales.'),
    ('Colores', 'En la pestaña Colores están los de la marca para gráficos y resaltados.'),
]
for i, (k, v) in enumerate(guide):
    r = 17 + i
    ws[f'B{r}'] = k; ws[f'B{r}'].font = font(10, True)
    ws[f'C{r}'] = v; ws[f'C{r}'].font = font(10, False, GRIS)
    ws.merge_cells(f'C{r}:E{r}')
    for c in 'BCDE':
        ws[f'{c}{r}'].border = Border(bottom=thin(NIEBLA))
    ws.row_dimensions[r].height = 22
ws['B17'].fill = FILL_IN

# ================= 2 · DATOS =================
ws = wb.create_sheet('Datos')
setup(ws, {'A': 3, 'B': 36, 'C': 11, 'D': 11, 'E': 16, 'F': 13, 'G': 16, 'H': 13, 'I': 3})
header_block(ws, 'Datos · ejemplo', 'Reservas por ruta', 'Tabla de datos con totales y porcentajes calculados. Datos de ejemplo.', 'H')
table_header(ws, 8, 'BCDEFGH', ['Ruta', 'Etapas', 'Km', 'Precio desde', 'Reservas', 'Ingresos', '% del total'],
             ['left', 'right', 'right', 'right', 'right', 'right', 'right'])
rutas = [
    ('Camino Francés desde Sarria', 5, 115, 695, 420),
    ('Camino Portugués desde Tui', 6, 118, 725, 260),
    ('Camino Inglés desde Ferrol', 5, 113, 680, 95),
    ('Camino del Norte desde Ribadeo', 8, 193, 1050, 70),
    ('Camino Primitivo desde Lugo', 5, 100, 710, 55),
]
first, last = 9, 9 + len(rutas) - 1
for i, (n, et, km, pr, res) in enumerate(rutas):
    r = first + i
    ws.row_dimensions[r].height = 22
    ws[f'B{r}'] = n; body_cell(ws[f'B{r}'], fill=FILL_IN)
    ws[f'C{r}'] = et; body_cell(ws[f'C{r}'], INT, 'right', FILL_IN)
    ws[f'D{r}'] = km; body_cell(ws[f'D{r}'], INT, 'right', FILL_IN)
    ws[f'E{r}'] = pr; body_cell(ws[f'E{r}'], EUR, 'right', FILL_IN)
    ws[f'F{r}'] = res; body_cell(ws[f'F{r}'], INT, 'right', FILL_IN)
    ws[f'G{r}'] = f'=E{r}*F{r}'; body_cell(ws[f'G{r}'], EUR, 'right')
    ws[f'H{r}'] = f'=IF($G${last + 1}=0,0,G{r}/$G${last + 1})'; body_cell(ws[f'H{r}'], PCT, 'right')
t = last + 1
ws.row_dimensions[t].height = 24
ws[f'B{t}'] = 'Total'; total_cell(ws[f'B{t}'])
for c in 'CDE':
    total_cell(ws[f'{c}{t}'])
ws[f'F{t}'] = f'=SUM(F{first}:F{last})'; total_cell(ws[f'F{t}'], INT, 'right')
ws[f'G{t}'] = f'=SUM(G{first}:G{last})'; total_cell(ws[f'G{t}'], EUR, 'right')
ws[f'H{t}'] = f'=SUM(H{first}:H{last})'; total_cell(ws[f'H{t}'], PCT, 'right')
ws[f'B{t + 2}'] = 'Fuente: datos de ejemplo. Precios por persona; ingresos = precio × reservas.'
ws[f'B{t + 2}'].font = font(8, False, GRIS)
ws.freeze_panes = 'C9'
ws.auto_filter.ref = f'B8:H{last}'

# ================= 3 · PRESUPUESTO =================
ws = wb.create_sheet('Presupuesto')
setup(ws, {'A': 3, 'B': 44, 'C': 12, 'D': 15, 'E': 16, 'F': 3}, landscape=False)
header_block(ws, 'Presupuesto', 'Presupuesto', 'Conceptos, cantidades e importes. Las celdas en verde claro se rellenan.', 'E')
for i, (k, v) in enumerate([('Cliente', 'Nombre Apellido'), ('Presupuesto n.º', '0000'), ('Fecha', 'dd/mm/aaaa'), ('Válido hasta', 'dd/mm/aaaa')]):
    c = 'BCDE'[i]
    ws[f'{c}8'] = k.upper(); ws[f'{c}8'].font = font(7, True, GRIS); ws[f'{c}8'].border = Border(top=medium(GRAFITO))
    ws[f'{c}9'] = v; ws[f'{c}9'].font = font(10); ws[f'{c}9'].fill = FILL_IN
ws.row_dimensions[10].height = 18
table_header(ws, 11, 'BCDE', ['Concepto', 'Cantidad', 'Precio', 'Importe'], ['left', 'right', 'right', 'right'])
conceptos = [
    ('Camino Francés desde Sarria · 7 días', 2, 695),
    ('Suplemento habitación individual', 1, 180),
    ('Cenas en ruta (5 noches)', 2, 120),
    ('Seguro de cancelación', 2, 35),
    ('', None, None), ('', None, None),
]
pf, pl = 12, 12 + len(conceptos) - 1
for i, (n, q, p) in enumerate(conceptos):
    r = pf + i
    ws.row_dimensions[r].height = 22
    ws[f'B{r}'] = n or None; body_cell(ws[f'B{r}'], fill=FILL_IN)
    ws[f'C{r}'] = q; body_cell(ws[f'C{r}'], INT, 'right', FILL_IN)
    ws[f'D{r}'] = p; body_cell(ws[f'D{r}'], EUR, 'right', FILL_IN)
    ws[f'E{r}'] = f'=IF(OR(C{r}="",D{r}=""),"",C{r}*D{r})'; body_cell(ws[f'E{r}'], EUR, 'right')
r = pl + 1
ws[f'D{r}'] = 'Subtotal'; ws[f'D{r}'].font = font(10); ws[f'D{r}'].alignment = Alignment(horizontal='right')
ws[f'E{r}'] = f'=SUM(E{pf}:E{pl})'; body_cell(ws[f'E{r}'], EUR, 'right')
ws[f'D{r + 1}'] = 'Descuento'; ws[f'D{r + 1}'].font = font(10); ws[f'D{r + 1}'].alignment = Alignment(horizontal='right')
ws[f'C{r + 1}'] = 0; body_cell(ws[f'C{r + 1}'], PCT, 'right', FILL_IN)
ws[f'C{r + 1}'].comment = Comment('Porcentaje de descuento sobre el subtotal (por ejemplo, 5 %). Déjalo en 0 si no hay descuento.', 'Santiago Ways')
ws[f'E{r + 1}'] = f'=-E{r}*C{r + 1}'; body_cell(ws[f'E{r + 1}'], EUR, 'right')
ws[f'D{r + 2}'] = 'Total'; total_cell(ws[f'D{r + 2}'], None, 'right')
ws[f'E{r + 2}'] = f'=E{r}+E{r + 1}'; total_cell(ws[f'E{r + 2}'], EUR, 'right')
for c in 'BC':
    total_cell(ws[f'{c}{r + 2}'])
ws.row_dimensions[r + 2].height = 24
ws[f'B{r + 4}'] = 'CONDICIONES'; ws[f'B{r + 4}'].font = font(8, True, SENDERO)
for i, txt in enumerate(['Precios de ejemplo, por persona salvo indicación.', 'Lorem ipsum dolor sit amet, consectetur adipiscing elit.', 'Integer posuere erat a ante venenatis dapibus.']):
    ws[f'B{r + 5 + i}'] = '•  ' + txt; ws[f'B{r + 5 + i}'].font = font(9, False, GRIS)

# ================= 4 · PLAN =================
ws = wb.create_sheet('Plan')
meses = ['Ene', 'Feb', 'Mar', 'Abr', 'May', 'Jun', 'Jul', 'Ago', 'Sep', 'Oct', 'Nov', 'Dic']
mcols = [chr(ord('G') + i) for i in range(12)]  # G..R
widths = {'A': 3, 'B': 32, 'C': 18, 'D': 9, 'E': 9, 'F': 12}
widths.update({c: 5.5 for c in mcols}); widths['S'] = 3
setup(ws, widths)
header_block(ws, 'Plan de trabajo', 'Plan de trabajo', 'Escribe el mes de inicio y de fin (1 a 12): las barras se pintan solas. Marca «Sí» en Hito para verlo en verde sendero.', 'R')
table_header(ws, 8, ['B', 'C', 'D', 'E', 'F'] + mcols, ['Tarea', 'Responsable', 'Inicio', 'Fin', 'Hito'] + meses,
             ['left', 'left', 'center', 'center', 'center'] + ['center'] * 12)
tareas = [
    ('Definición de la campaña', 'Nombre Apellido', 1, 2, 'No'),
    ('Producción de contenidos', 'Nombre Apellido', 2, 4, 'No'),
    ('Lanzamiento', 'Nombre Apellido', 4, 4, 'Sí'),
    ('Campaña en redes', 'Nombre Apellido', 4, 9, 'No'),
    ('Medición y ajustes', 'Nombre Apellido', 6, 11, 'No'),
    ('Informe final', 'Nombre Apellido', 12, 12, 'Sí'),
    ('', '', None, None, ''), ('', '', None, None, ''),
]
gf, gl = 9, 9 + len(tareas) - 1
for i, (t_, resp, a, b, h) in enumerate(tareas):
    r = gf + i
    ws.row_dimensions[r].height = 22
    ws[f'B{r}'] = t_ or None; body_cell(ws[f'B{r}'], fill=FILL_IN)
    ws[f'C{r}'] = resp or None; body_cell(ws[f'C{r}'], fill=FILL_IN)
    ws[f'D{r}'] = a; body_cell(ws[f'D{r}'], '0', 'center', FILL_IN)
    ws[f'E{r}'] = b; body_cell(ws[f'E{r}'], '0', 'center', FILL_IN)
    ws[f'F{r}'] = h or None; body_cell(ws[f'F{r}'], None, 'center', FILL_IN)
    for c in mcols:
        cell = ws[f'{c}{r}']
        cell.border = Border(bottom=thin(NIEBLA), left=thin(BLANCO), right=thin(BLANCO))
dv = DataValidation(type='list', formula1='"Sí,No"', allow_blank=True)
ws.add_data_validation(dv); dv.add(f'F{gf}:F{gl}')
dvm = DataValidation(type='whole', operator='between', formula1='1', formula2='12', allow_blank=True)
ws.add_data_validation(dvm); dvm.add(f'D{gf}:E{gl}')
rng = f'G{gf}:R{gl}'
# la columna G es el mes 1: COLUMN()-6 da el número de mes
ws.conditional_formatting.add(rng, FormulaRule(formula=[f'AND($D{gf}<>"",$E{gf}<>"",COLUMN()-6>=$D{gf},COLUMN()-6<=$E{gf},$F{gf}="Sí")'], fill=PatternFill('solid', fgColor=SENDERO, bgColor=SENDERO), stopIfTrue=True))
ws.conditional_formatting.add(rng, FormulaRule(formula=[f'AND($D{gf}<>"",$E{gf}<>"",COLUMN()-6>=$D{gf},COLUMN()-6<=$E{gf})'], fill=PatternFill('solid', fgColor=WAYS, bgColor=WAYS)))
ws.freeze_panes = 'D9'

# ================= 5 · RESUMEN =================
ws = wb.create_sheet('Resumen')
setup(ws, {'A': 3, 'B': 24, 'C': 3, 'D': 24, 'E': 3, 'F': 24, 'G': 3, 'H': 24, 'I': 3})
header_block(ws, 'Resumen', 'Cifras clave', 'Se calcula a partir de la pestaña Datos. No hay que rellenar nada.', 'H')
kp = [
    ('B', 'Reservas', f"=Datos!F{t}", INT),
    ('D', 'Ingresos', f"=Datos!G{t}", EUR),
    ('F', 'Ticket medio', f"=IF(Datos!F{t}=0,0,Datos!G{t}/Datos!F{t})", EUR),
    ('H', 'Ruta con más reservas', f"=INDEX(Datos!B{first}:B{last},MATCH(MAX(Datos!F{first}:F{last}),Datos!F{first}:F{last},0))", None),
]
for c, lab, fml, fmt in kp:
    ws[f'{c}8'].border = Border(top=Side(style='thick', color=WAYS))
    ws[f'{c}9'] = fml; ws[f'{c}9'].font = font(18 if fmt else 11, True); ws[f'{c}9'].alignment = Alignment(horizontal='left', vertical='bottom', wrap_text=True)
    if fmt:
        ws[f'{c}9'].number_format = fmt
    ws[f'{c}10'] = lab; ws[f'{c}10'].font = font(9, False, GRIS)
ws.row_dimensions[9].height = 34
ws['B12'] = 'INGRESOS POR RUTA'; ws['B12'].font = font(8, True, SENDERO)
ch = BarChart()
ch.type = 'bar'
ch.style = 10
ch.title = None
ch.legend = None
ch.y_axis.majorGridlines = None
ch.y_axis.delete = True
ch.x_axis.scaling.orientation = 'maxMin'
ch.x_axis.delete = False
data = Reference(wb['Datos'], min_col=7, min_row=first, max_row=last)
cats = Reference(wb['Datos'], min_col=2, min_row=first, max_row=last)
ch.add_data(data, titles_from_data=False)
ch.set_categories(cats)
s0 = ch.series[0]
s0.graphicalProperties.solidFill = WAYS
s0.graphicalProperties.line.noFill = True
ch.gapWidth = 60
ch.dataLabels = DataLabelList(); ch.dataLabels.showVal = True; ch.dataLabels.showCatName = False; ch.dataLabels.showSerName = False; ch.dataLabels.showLegendKey = False; ch.dataLabels.showPercent = False; ch.dataLabels.numFmt = '#,##0 "€"'
ch.height, ch.width = 8.5, 22
from openpyxl.chart.text import RichText
from openpyxl.drawing.text import Paragraph as DPara, ParagraphProperties, CharacterProperties, Font as DFont
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.drawing.line import LineProperties
def chart_text(bold=False):
    cp = CharacterProperties(latin=DFont(typeface=FONT), sz=900, b=bold, solidFill=GRAFITO)
    return RichText(p=[DPara(pPr=ParagraphProperties(defRPr=cp), endParaRPr=cp)])
ch.x_axis.txPr = chart_text()
ch.dataLabels.txPr = chart_text(True)
ch.x_axis.graphicalProperties = GraphicalProperties(ln=LineProperties(solidFill=NIEBLA))
try:
    ch.graphical_properties = GraphicalProperties(ln=LineProperties(noFill=True))
except Exception:
    pass
ws.add_chart(ch, 'B13')

# ================= 6 · COLORES =================
ws = wb.create_sheet('Colores')
setup(ws, {'A': 3, 'B': 8, 'C': 20, 'D': 12, 'E': 52, 'F': 3})
header_block(ws, 'Marca', 'Colores de la marca', 'Úsalos en gráficos, cabeceras y resaltados. El verde Ways nunca como color de texto sobre blanco.', 'E')
table_header(ws, 8, 'BCDE', ['', 'Color', 'Hex', 'Uso'], ['left'] * 4)
pal = [
    (BLANCO, 'Blanco', 'El fondo principal.'), (WAYS, 'Verde Ways', 'Barras de gráficos, filetes y acentos.'),
    (SENDERO, 'Verde sendero', 'Resaltar un dato o una serie; antetítulos.'), (GRAFITO, 'Grafito', 'Texto principal y cabeceras de tabla.'),
    (GRIS, 'Gris', 'Textos secundarios y notas.'), (NIEBLA, 'Niebla', 'Líneas finas entre filas.'),
    (BROTE, 'Brote', 'Celdas para rellenar.'),
]
for i, (hx, n, u) in enumerate(pal):
    r = 9 + i
    ws.row_dimensions[r].height = 24
    ws[f'B{r}'].fill = PatternFill('solid', fgColor=hx)
    ws[f'B{r}'].border = Border(top=thin(NIEBLA), bottom=thin(NIEBLA), left=thin(NIEBLA), right=thin(NIEBLA))
    ws[f'C{r}'] = n; body_cell(ws[f'C{r}'], bold=True)
    ws[f'D{r}'] = '#' + hx; body_cell(ws[f'D{r}'])
    ws[f'E{r}'] = u; body_cell(ws[f'E{r}'])
    ws[f'E{r}'].font = font(10, False, GRIS)

wb.active = 0
# Excel y Hojas de cálculo de Google calculan todas las fórmulas al abrir
from openpyxl.workbook.properties import CalcProperties
wb.calculation = CalcProperties(fullCalcOnLoad=True)
wb.save(OUT)
print('ok', OUT)
